"""
report.py — builds the daily usage/feedback Excel report and emails it.

Deliberately isolated from app.py: this module is the ONLY place an email
ever gets sent from, and the only place that reads the email API key. It is
never imported by, or wired into, any user-facing request path (/api/generate,
/api/comment, etc.) — only the gated admin trigger in app.py calls into it,
so a bug or abuse of a public endpoint can't cause a send.
"""

import os
import io
import base64
from datetime import datetime, timedelta, timezone

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

import db

# Resend (not Gmail SMTP) per explicit decision — a transactional email API is
# more reliable from a server IP than personal SMTP, which mail providers
# throttle/flag. The API key and recipient are read only from the
# environment: never hardcoded in a way that could leak via source control,
# and never accepted from a request — see app.py's admin route, which takes
# no email-related input at all.
RESEND_API_KEY = os.getenv("RESEND_API_KEY")
REPORT_SENDER_EMAIL = os.getenv("REPORT_SENDER_EMAIL", "onboarding@resend.dev")
# Hardcoded fallback is the team's own known inbox, not a guess — still
# overridable via env var per-deployment, but a request can never set this.
REPORT_RECIPIENT_EMAIL = os.getenv("REPORT_RECIPIENT_EMAIL", "bluemachine.tools@apna.co")

LANGUAGE_DISPLAY = {
    "HI": "Hindi", "EN": "English", "KN": "Kannada", "TA": "Tamil",
    "ML": "Malayalam", "GU": "Gujarati", "MR": "Marathi", "TE": "Telugu",
    "OD": "Odia", "BN": "Bengali",
}

# Excel cells cap at 32,767 characters — stay comfortably under that so a
# very long prompt can never corrupt the workbook.
_CELL_CHAR_LIMIT = 30000
# Above this combined length, showing both the raw input AND the business
# logic in one cell stops being readable — fall back to business logic only,
# since that's the more useful of the two for reviewing what was generated.
_COMBINED_INLINE_LIMIT = 4000


def _truncate(text: str, limit: int = _CELL_CHAR_LIMIT) -> str:
    text = text or ""
    if len(text) > limit:
        return text[:limit] + "... [truncated]"
    return text


def _format_input_column(raw_prompt: str, business_logic: str) -> str:
    raw_prompt = raw_prompt or ""
    business_logic = business_logic or ""
    if not business_logic.strip():
        return _truncate(raw_prompt)
    combined = f"INPUT:\n{raw_prompt}\n\nBUSINESS LOGIC:\n{business_logic}"
    if len(combined) <= _COMBINED_INLINE_LIMIT:
        return _truncate(combined)
    return _truncate(f"BUSINESS LOGIC:\n{business_logic}")


def _format_languages(languages_requested: str) -> str:
    codes = [c.strip() for c in (languages_requested or "").split(",") if c.strip()]
    return ", ".join(LANGUAGE_DISPLAY.get(c, c) for c in codes)


def _format_comments(comments: list) -> str:
    if not comments:
        return ""
    parts = [f"{(c.get('commenter_name') or 'Anonymous')}: {c.get('feedback_text') or ''}" for c in comments]
    return _truncate(" | ".join(parts))


def build_report_workbook(generations: list) -> bytes:
    """Builds the report as an in-memory .xlsx file (same openpyxl approach
    used for Prompt_Studio_Evaluation.xlsx earlier this session) and returns
    its raw bytes, ready to attach to an email or write to disk."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Report"

    headers = ["Timestamp", "User Name", "Input / Business Prompt",
               "Languages Generated", "Model Used", "Comments / Feedback"]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for gen in generations:
        ws.append([
            gen.get("timestamp", ""),
            gen.get("user_name", ""),
            _format_input_column(gen.get("raw_prompt"), gen.get("business_logic")),
            _format_languages(gen.get("languages_requested")),
            gen.get("model_used", ""),
            _format_comments(gen.get("comments")),
        ])

    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 1)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for col_idx, width in enumerate([20, 16, 60, 24, 18, 50], start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def send_report_email(xlsx_bytes: bytes, report_date_label: str) -> dict:
    if not RESEND_API_KEY:
        raise RuntimeError("RESEND_API_KEY is not configured — cannot send the daily report email.")

    payload = {
        "from": REPORT_SENDER_EMAIL,
        "to": [REPORT_RECIPIENT_EMAIL],
        "subject": f"Prompt Studio — Daily Usage Report ({report_date_label})",
        "text": f"Attached: Prompt Studio generation and feedback log for {report_date_label}.",
        "attachments": [{
            "filename": f"prompt_studio_report_{report_date_label}.xlsx",
            "content": base64.b64encode(xlsx_bytes).decode("ascii"),
        }],
    }
    resp = requests.post(
        "https://api.resend.com/emails",
        headers={"Authorization": f"Bearer {RESEND_API_KEY}", "Content-Type": "application/json"},
        json=payload,
        timeout=30,
    )
    resp.raise_for_status()
    return resp.json()


def generate_and_send_daily_report(hours: int = 24) -> dict:
    """Pulls every generation (+ its comments) from the last `hours` hours,
    builds the workbook, and emails it. Returns a small summary dict — never
    the raw workbook bytes or anything email-credential-related, since this
    return value flows back through the admin endpoint's JSON response."""
    cutoff = (datetime.now(timezone.utc) - timedelta(hours=hours)).isoformat()
    generations = db.get_generations_since(cutoff)
    xlsx_bytes = build_report_workbook(generations)
    report_date_label = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    email_result = send_report_email(xlsx_bytes, report_date_label)
    return {"rows": len(generations), "date": report_date_label, "email_id": email_result.get("id")}
